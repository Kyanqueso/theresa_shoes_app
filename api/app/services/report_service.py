"""The downloadable business report — one Excel workbook, one tab per topic.

Built on the server rather than in the browser so it covers every order (the list endpoints
page at 100 rows) and so its totals are worked out by the same rules as the Analytics page:
archived orders are left out entirely, and money received is dated by when each installment
arrived. If the report and the dashboard ever disagree, one of them is wrong.

Tabs:
  Summary      headline totals, then a month-by-month table
  Orders       one row per order, every spec the shop needs to make it
  Payments     one row per order, each installment with its date
  Uncollected  only orders still owing money or not yet delivered
  Companies    one row per company, with what it has ordered, paid and still owes
"""

from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, joinedload, selectinload

from app.config.timezone import business_now, to_business
from app.db.models import Company, CompanyStatus, Order, OrderStatus, Payment

ZERO = Decimal("0")

# Brand colours, matching the app.
MAROON = "3A0D1F"
CREAM = "EFE9E1"

MONEY = '"₱"#,##0.00'
DATE = "mmm d, yyyy"  # "Sep 5, 2026" — the same format the app shows
COUNT = "#,##0"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=MAROON)
TITLE_FONT = Font(bold=True, size=16, color=MAROON)
SECTION_FONT = Font(bold=True, size=12, color=MAROON)
BOLD = Font(bold=True)
MUTED = Font(italic=True, color="6B5D61")
RULE = Border(bottom=Side(style="thin", color="A67D44"))

# Same four states, labels and colours as the Payments page (web/src/lib/paymentStatus.js).
FULFILMENT = {
    "fulfilled": ("Fulfilled", "C6EFCE"),
    "unpaid": ("Unpaid", "FED7AA"),
    "undelivered": ("Undelivered", "BFDBFE"),
    "unpaid_undelivered": ("Unpaid & Undelivered", "FFFFFF"),
}


# ---------------------------------------------------------------- reading the data


def _option_name(option) -> str:
    return option.name if option is not None else ""


def _extra(flag: bool, option) -> str:
    """Buckle / flatform / slingback: "No", "Yes", or "Yes – <which one>"."""
    if not flag:
        return "No"
    return f"Yes – {option.name}" if option is not None else "Yes"


def _notes_text(blocks) -> str:
    """The typed parts of an order's notes. Photos and drawings can't go in a spreadsheet, so
    they're counted instead — enough to know there's something to look at in the app."""
    if not isinstance(blocks, list):
        return ""
    texts = [
        block["value"].strip()
        for block in blocks
        if isinstance(block, dict) and block.get("type") == "text" and str(block.get("value") or "").strip()
    ]
    pictures = sum(1 for block in blocks if isinstance(block, dict) and block.get("type") in ("photo", "drawing"))
    if pictures:
        texts.append(f"[{pictures} photo{'s' if pictures != 1 else ''}/drawing{'s' if pictures != 1 else ''} in the app]")
    return " / ".join(texts)


def _local_date(value: datetime | None) -> date | None:
    """A stored UTC timestamp as the calendar day it was in the shop."""
    return to_business(value).date() if value is not None else None


def _order_number(n: int) -> str:
    # Mirrors formatOrderNumber() in the web app.
    return f"ORDER-{n:03d}"


class _Row:
    """Everything the report says about one order, worked out once."""

    def __init__(self, order: Order):
        payment: Payment | None = order.payment
        self.order = order
        self.number = _order_number(order.order_number)
        self.ordered_on = _local_date(order.created_at)
        self.company = order.company.name if order.company else ""
        self.client = order.client_name or ""
        self.contact = order.contact_number or ""
        self.model = order.shoe.name if order.shoe else (order.custom_model_name or "")
        self.quantity = order.quantity or 0
        self.unit_price = Decimal(order.unit_price or 0)
        self.order_total = self.unit_price * self.quantity

        if payment is not None:
            self.total = Decimal(payment.total_amount or 0)
            self.installments = [
                (Decimal(payment.first_payment or 0), payment.first_payment_date),
                (Decimal(payment.second_payment or 0), payment.second_payment_date),
                (Decimal(payment.third_payment or 0), payment.third_payment_date),
            ]
            self.balance = Decimal(payment.balance or 0)
            self.cleared_on = payment.balance_cleared_date
            self.delivered_on = payment.date_delivered
        else:
            # Every order is created with a payment row; this only guards against a bad one.
            self.total = self.order_total
            self.installments = [(ZERO, None)] * 3
            self.balance = self.order_total
            self.cleared_on = None
            self.delivered_on = None

        self.paid = sum((amount for amount, _ in self.installments), ZERO)
        is_paid = self.balance <= 0
        is_delivered = self.delivered_on is not None
        if is_paid and is_delivered:
            self.fulfilment = "fulfilled"
        elif not is_paid and not is_delivered:
            self.fulfilment = "unpaid_undelivered"
        elif not is_paid:
            self.fulfilment = "unpaid"
        else:
            self.fulfilment = "undelivered"

    def received(self):
        """(amount, day) for every installment actually received. Installments recorded before
        per-installment dates existed fall back to the order's date — the same rule the
        Analytics page uses, so the monthly figures match."""
        return [(amount, day or self.ordered_on) for amount, day in self.installments if amount > 0]


def _load_rows(db: Session) -> list[_Row]:
    orders = (
        db.query(Order)
        .filter(Order.status != OrderStatus.archived)
        .options(
            joinedload(Order.company),
            joinedload(Order.shoe),
            joinedload(Order.payment),
            selectinload(Order.material),
            selectinload(Order.mold_type),
            selectinload(Order.heel_type),
            selectinload(Order.buckle),
            selectinload(Order.flatform),
            selectinload(Order.slingback),
        )
        .order_by(Order.order_number)
        .all()
    )
    return [_Row(order) for order in orders]


# ---------------------------------------------------------------- writing the workbook


def _table(ws, columns, rows, *, freeze="B2"):
    """A plain data tab: one header row, then the rows, filterable and frozen.

    `columns` is a list of (heading, number_format or None, width or None). Widths left as
    None are sized to the longest value in the column.
    """
    ws.append([heading for heading, _, _ in columns])
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for values in rows:
        ws.append(values)

    for index, (heading, number_format, width) in enumerate(columns, start=1):
        letter = get_column_letter(index)
        if number_format:
            for (cell,) in ws.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = number_format
        if width is None:
            longest = max(
                [len(heading)] + [len(_display(row[index - 1], number_format)) for row in rows],
            )
            width = min(max(longest + 2, 8), 50)
        ws.column_dimensions[letter].width = width

    last_column = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last_column}{max(len(rows) + 1, 2)}"
    ws.freeze_panes = freeze


def _display(value, number_format) -> str:
    """Roughly how wide a value will look once formatted — only used to size columns."""
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return "Sep 30, 2026"
    if number_format == MONEY:
        return f"₱{value:,.2f}"
    return str(value)


def _colour_status(ws, column_index: int, rows: list[_Row]):
    for offset, row in enumerate(rows, start=2):
        _, colour = FULFILMENT[row.fulfilment]
        ws.cell(row=offset, column=column_index).fill = PatternFill("solid", fgColor=colour)


def _money_or_blank(amount: Decimal):
    return amount if amount else None


def _summary_sheet(ws, rows: list[_Row], generated_at: datetime):
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    for letter in "BCD":
        ws.column_dimensions[letter].width = 20

    ws["A1"] = "Theresa Shoes — Business Report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {_stamp(generated_at)} (Philippine time)"
    ws["A2"].font = MUTED
    ws["A3"] = "Covers every order ever placed. Archived orders are left out, the same as on the Analytics page."
    ws["A3"].font = MUTED

    current = [row for row in rows if row.order.status == OrderStatus.current]
    completed = [row for row in rows if row.order.status == OrderStatus.completed]
    owing = [row for row in rows if row.balance > 0]
    undelivered = [row for row in rows if row.delivered_on is None]
    companies = {row.company for row in rows if row.company}

    figures = [
        ("Orders", len(rows), COUNT),
        ("   Current", len(current), COUNT),
        ("   Completed", len(completed), COUNT),
        ("Pairs ordered", sum(row.quantity for row in rows), COUNT),
        ("Companies with orders", len(companies), COUNT),
        (None, None, None),
        ("Value of all orders", sum((row.total for row in rows), ZERO), MONEY),
        ("Payments received", sum((row.paid for row in rows), ZERO), MONEY),
        ("Still owed", sum((row.balance for row in owing), ZERO), MONEY),
        ("Orders still owing", len(owing), COUNT),
        ("Orders not yet delivered", len(undelivered), COUNT),
    ]

    ws["A5"] = "At a glance"
    ws["A5"].font = SECTION_FONT
    ws["A5"].border = RULE
    ws["B5"].border = RULE
    line = 6
    for label, value, number_format in figures:
        if label is not None:
            ws.cell(row=line, column=1, value=label)
            cell = ws.cell(row=line, column=2, value=value)
            cell.number_format = number_format
            if not label.startswith(" "):
                ws.cell(row=line, column=1).font = BOLD
                cell.font = BOLD
        line += 1

    # Month by month: orders are counted in the month they were placed; money in the month it
    # arrived. Only months with something in them are listed — a backdated 1998 order would
    # otherwise drag in three hundred empty rows.
    placed = defaultdict(lambda: [0, ZERO])
    received = defaultdict(lambda: ZERO)
    for row in rows:
        if row.ordered_on is not None:
            key = (row.ordered_on.year, row.ordered_on.month)
            placed[key][0] += 1
            placed[key][1] += row.total
        for amount, day in row.received():
            if day is not None:
                received[(day.year, day.month)] += amount

    line += 1
    ws.cell(row=line, column=1, value="Month by month").font = SECTION_FONT
    for column in range(1, 5):
        ws.cell(row=line, column=column).border = RULE
    line += 1
    headings = ["Month", "Orders placed", "Value of orders placed", "Payments received"]
    for column, heading in enumerate(headings, start=1):
        cell = ws.cell(row=line, column=column, value=heading)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    line += 1

    months = sorted(set(placed) | set(received), reverse=True)
    if not months:
        ws.cell(row=line, column=1, value="No orders yet.").font = MUTED
    for year, month in months:
        ws.cell(row=line, column=1, value=date(year, month, 1)).number_format = "mmmm yyyy"
        ws.cell(row=line, column=1).alignment = Alignment(horizontal="left")
        count, value = placed.get((year, month), (0, ZERO))
        ws.cell(row=line, column=2, value=count).number_format = COUNT
        ws.cell(row=line, column=3, value=value).number_format = MONEY
        ws.cell(row=line, column=4, value=received.get((year, month), ZERO)).number_format = MONEY
        if (year, month) == (generated_at.year, generated_at.month):
            for column in range(1, 5):
                ws.cell(row=line, column=column).fill = PatternFill("solid", fgColor=CREAM)
        line += 1

    line += 1
    ws.cell(
        row=line,
        column=1,
        value="Payments received counts each installment in the month it was paid. Installments "
        "recorded before payment dates were kept are counted in the month of the order.",
    ).font = MUTED


def _stamp(moment: datetime) -> str:
    # strftime's no-padding flags differ between Windows and Linux, so build it by hand.
    hour = moment.hour % 12 or 12
    return f"{moment.strftime('%b')} {moment.day}, {moment.year} at {hour}:{moment.minute:02d} {'AM' if moment.hour < 12 else 'PM'}"


def _orders_sheet(ws, rows: list[_Row]):
    ws.title = "Orders"
    columns = [
        ("Order ID", None, 12),
        ("Order Date", DATE, 14),
        ("Status", None, 11),
        ("Company", None, None),
        ("Client", None, None),
        ("Contact #", None, 15),
        ("Model", None, None),
        ("Material", None, None),
        ("Color / Code", None, None),
        ("Mold Type", None, None),
        ("Heel Type", None, None),
        ("Heel Size", None, 10),
        ("Size", None, 8),
        ("Buckle", None, None),
        ("Flatform", None, None),
        ("Slingback", None, None),
        ("Qty", COUNT, 7),
        ("Unit Price", MONEY, 13),
        ("Total", MONEY, 14),
        ("Completed On", DATE, 14),
        ("Notes", None, 50),
    ]
    data = []
    for row in rows:
        order = row.order
        data.append(
            [
                row.number,
                row.ordered_on,
                order.status.value.title(),
                row.company,
                row.client,
                row.contact,  # kept as text so the leading 0 of 09… survives
                row.model,
                _option_name(order.material),
                order.color_code or "",
                _option_name(order.mold_type),
                _option_name(order.heel_type),
                order.heel_size,
                order.size,
                _extra(order.with_buckle, order.buckle),
                _extra(order.with_flatform, order.flatform),
                _extra(order.with_slingback, order.slingback),
                row.quantity,
                row.unit_price,
                row.order_total,
                _local_date(order.completed_at) if order.status == OrderStatus.completed else None,
                _notes_text(order.notes_blocks),
            ]
        )
    _table(ws, columns, data)
    for (cell,) in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        cell.number_format = "@"
    for (cell,) in ws.iter_rows(min_row=2, min_col=21, max_col=21):
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _payments_sheet(ws, rows: list[_Row]):
    ws.title = "Payments"
    columns = [
        ("Order ID", None, 12),
        ("Order Date", DATE, 14),
        ("Company", None, None),
        ("Client", None, None),
        ("Total", MONEY, 14),
        ("1st Payment", MONEY, 13),
        ("1st Paid On", DATE, 14),
        ("2nd Payment", MONEY, 13),
        ("2nd Paid On", DATE, 14),
        ("3rd Payment", MONEY, 13),
        ("3rd Paid On", DATE, 14),
        ("Total Paid", MONEY, 14),
        ("Balance", MONEY, 14),
        ("Fully Paid On", DATE, 14),
        ("Delivered On", DATE, 14),
        ("Status", None, 22),
    ]
    data = []
    for row in rows:
        (first, first_on), (second, second_on), (third, third_on) = row.installments
        data.append(
            [
                row.number,
                row.ordered_on,
                row.company,
                row.client,
                row.total,
                _money_or_blank(first),
                first_on if first else None,
                _money_or_blank(second),
                second_on if second else None,
                _money_or_blank(third),
                third_on if third else None,
                row.paid,
                row.balance,
                row.cleared_on,
                row.delivered_on,
                FULFILMENT[row.fulfilment][0],
            ]
        )
    _table(ws, columns, data)
    _colour_status(ws, len(columns), rows)


def _uncollected_sheet(ws, rows: list[_Row]):
    ws.title = "Uncollected"
    # Oldest first: the longer something has been outstanding, the sooner it needs chasing.
    pending = sorted(
        (row for row in rows if row.fulfilment != "fulfilled"),
        key=lambda row: (row.ordered_on or date.max, row.order.order_number),
    )
    columns = [
        ("Order ID", None, 12),
        ("Order Date", DATE, 14),
        ("Company", None, None),
        ("Client", None, None),
        ("Contact #", None, 15),
        ("Model", None, None),
        ("Qty", COUNT, 7),
        ("Total", MONEY, 14),
        ("Paid So Far", MONEY, 14),
        ("Balance", MONEY, 14),
        ("Delivered On", DATE, 14),
        ("Status", None, 22),
    ]
    data = [
        [
            row.number,
            row.ordered_on,
            row.company,
            row.client,
            row.contact,
            row.model,
            row.quantity,
            row.total,
            row.paid,
            row.balance,
            row.delivered_on,
            FULFILMENT[row.fulfilment][0],
        ]
        for row in pending
    ]
    _table(ws, columns, data)
    for (cell,) in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        cell.number_format = "@"
    _colour_status(ws, len(columns), pending)
    if not pending:
        ws.cell(row=2, column=1, value="Nothing outstanding — every order is paid and delivered.").font = MUTED


def _companies_sheet(ws, rows: list[_Row], companies: list[Company]):
    ws.title = "Companies"
    by_company = defaultdict(list)
    for row in rows:
        by_company[row.order.company_id].append(row)

    listed = []
    for company in companies:
        company_rows = by_company.get(company.id, [])
        # Archived companies only have archived orders (archiving cascades), so they'd be all
        # zeros — list them only in the odd case one still has an order that counts.
        if company.status == CompanyStatus.archive and not company_rows:
            continue
        listed.append((company.name, "Archived" if company.status == CompanyStatus.archive else "Active", company_rows))
    listed.sort(key=lambda item: item[0].lower())
    if by_company.get(None):
        listed.append(("(No company)", "", by_company[None]))

    columns = [
        ("Company", None, None),
        ("Status", None, 10),
        ("Orders", COUNT, 9),
        ("Pairs", COUNT, 9),
        ("Value of Orders", MONEY, 16),
        ("Paid", MONEY, 16),
        ("Still Owed", MONEY, 16),
        ("Not Yet Delivered", COUNT, 12),
        ("Latest Order", DATE, 14),
    ]
    data = []
    for name, status, company_rows in listed:
        dates = [row.ordered_on for row in company_rows if row.ordered_on is not None]
        data.append(
            [
                name,
                status,
                len(company_rows),
                sum(row.quantity for row in company_rows),
                sum((row.total for row in company_rows), ZERO),
                sum((row.paid for row in company_rows), ZERO),
                sum((row.balance for row in company_rows if row.balance > 0), ZERO),
                sum(1 for row in company_rows if row.delivered_on is None),
                max(dates) if dates else None,
            ]
        )
    _table(ws, columns, data)


def build_report(db: Session) -> tuple[bytes, str]:
    """The finished workbook, and the filename to save it under."""
    generated_at = business_now()
    rows = _load_rows(db)
    companies = db.query(Company).all()

    workbook = Workbook()
    _summary_sheet(workbook.active, rows, generated_at)
    _orders_sheet(workbook.create_sheet(), rows)
    _payments_sheet(workbook.create_sheet(), rows)
    _uncollected_sheet(workbook.create_sheet(), rows)
    _companies_sheet(workbook.create_sheet(), rows, companies)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), f"theresa-shoes-report-{generated_at.date().isoformat()}.xlsx"
